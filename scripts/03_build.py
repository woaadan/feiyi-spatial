# -*- coding: utf-8 -*-
"""合并原表与地理编码结果，生成主表 / 点位表 / 质量报告。
输出 data/projects.csv, data/points.csv, data/qa_report.json
"""
import json, os, re, csv, collections
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GEO = os.path.join(ROOT, "data", "geo")
DATA = os.path.join(ROOT, "data")
SRC = r"C:\Users\11487\Desktop\数字人文\空间人文\国家级非物质文化遗产代表性项目名录_全量4234条.xlsx"

CRAFT = {
    "FRZX": "纺染织绣", "FSZZ": "服饰制作", "BZZZ": "编织扎制", "DKSZ": "雕刻塑造",
    "JJJZ": "家具建筑", "JSJG": "金属加工", "JZKH": "剪纸刻绘", "TCSZ": "陶瓷烧造",
    "WFZZ": "文房制作", "QQXS": "漆器髹饰", "YSZB": "印刷装裱", "SPZZ": "食品制作",
    "ZYPZ": "中药炮制", "QJZZ": "器具制作",
}
BATCH_MAP = {"一批": "第一批 2006", "二批": "第二批 2008", "三批": "第三批 2011",
             "四批": "第四批 2014", "五批": "第五批 2021"}


def load_polys():
    gj = json.load(open(os.path.join(GEO, "china_provinces.json"), encoding="utf-8"))
    polys = []
    for f in gj["features"]:
        g = f["geometry"]
        if not g:
            continue
        if g["type"] == "Polygon":
            rings = [g["coordinates"]]
        else:
            rings = g["coordinates"]
        polys.append((f["properties"]["name"], f["properties"]["adcode"], rings))
    return polys


def point_in_ring(x, y, ring):
    inside = False
    n = len(ring)
    j = n - 1
    for i in range(n):
        xi, yi = ring[i][0], ring[i][1]
        xj, yj = ring[j][0], ring[j][1]
        if ((yi > y) != (yj > y)) and (x < (xj - xi) * (y - yi) / (yj - yi + 1e-15) + xi):
            inside = not inside
        j = i
    return inside


def point_in_province(x, y, polys):
    hits = []
    for name, code, rings in polys:
        for poly in rings:
            if not poly or len(poly[0]) < 4:
                continue
            if point_in_ring(x, y, poly[0]):
                # 检查洞
                if not any(len(h) >= 4 and point_in_ring(x, y, h) for h in poly[1:]):
                    hits.append((name, code))
                    break
    return hits


def east_of_hu_line(lng, lat):
    """胡焕庸线：黑河(127.5,50.25)—腾冲(98.5,25.0)，返回是否在线东南侧。"""
    x1, y1, x2, y2 = 127.5, 50.25, 98.5, 25.0
    cross = (x2 - x1) * (lat - y1) - (y2 - y1) * (lng - x1)
    return cross > 0


def main():
    ws = openpyxl.load_workbook(SRC, read_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] is not None]
    places = json.load(open(os.path.join(GEO, "places2.json"), encoding="utf-8"))
    by_seq = collections.defaultdict(list)
    for p in places:
        by_seq[p["seq"]].append(p)

    polys = load_polys()
    projects, points, qa = [], [], []
    outside, nohit, dupname = 0, 0, 0
    for r in rows:
        seq, code, serial, name, cat, pub, batch, typ, raw, protect, link = r[:11]
        serial = str(serial or "")
        m = re.search(r"-([A-Z]{4})-", serial)
        dataset = "传统工艺振兴目录" if m else "国家级非遗名录"
        if m:
            cname = CRAFT.get(m.group(1), "传统工艺")
            cat = f"传统工艺·{cname}"
            batch_l = "2018振兴目录"
            typ_l = "振兴目录"
        else:
            batch_l = BATCH_MAP.get(batch, str(batch))
            typ_l = typ or ""
        segs = [p for p in by_seq.get(seq, []) if p.get("ok")]
        for p in segs:
            pts = {
                "seq": seq, "名称": name, "类别": cat or "未分类", "批次": batch_l,
                "类型": typ_l, "数据集": dataset, "申报地区": raw,
                "地点片段": p["raw"], "省": p["prov"], "省代码": p["prov_code"],
                "市": p.get("city") or "", "市代码": p.get("city_code") or "",
                "县区": p.get("county") or "", "县区代码": p.get("county_code") or "",
                "经度": round(p["lng"], 6), "纬度": round(p["lat"], 6),
                "定位精度": {3: "县级", 2: "市级", 1: "省级"}[p["precision"]],
                "地点类型": p.get("kind") or "行政区",
                "链接": link,
            }
            points.append(pts)

        if not segs:
            continue
        prim = sorted(segs, key=lambda p: (-p["precision"], p["lng"]))[0]
        proj = {
            "序号": seq, "项目编号": code or "", "名称": name, "类别": cat or "未分类",
            "批次": batch_l, "类型": typ_l, "数据集": dataset,
            "申报地区": raw, "保护单位": protect or "", "点位数": len(segs),
            "主点位_地点": prim["raw"], "省": prim["prov"], "市": prim.get("city") or "",
            "县区": prim.get("county") or "", "县区代码": prim.get("county_code") or "",
            "经度": round(prim["lng"], 6), "纬度": round(prim["lat"], 6),
            "最高精度": {3: "县级", 2: "市级", 1: "省级"}[prim["precision"]],
            "链接": link,
        }
        projects.append(proj)

    # 几何校验：点位是否落在其所属省级面内
    bad = []
    for p in points:
        hits = point_in_province(p["经度"], p["纬度"], polys)
        names = [h[0] for h in hits]
        if not hits:
            nohit += 1
            bad.append([p["seq"], p["名称"], p["省"], p["地点片段"], p["经度"], p["纬度"], "不在任何省面内"])
        elif p["省"] not in names:
            outside += 1
            bad.append([p["seq"], p["名称"], p["省"], p["地点片段"], p["经度"], p["纬度"],
                        "落在" + "/".join(names)])

    for p, d in zip(points, [None] * len(points)):
        p["胡焕庸线以东"] = "是" if east_of_hu_line(p["经度"], p["纬度"]) else "否"

    # 写 CSV
    def dump(path, recs, fields=None):
        if not recs:
            return
        fields = fields or list(recs[0].keys())
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader(); w.writerows(recs)

    dump(os.path.join(DATA, "points.csv"), points)
    dump(os.path.join(DATA, "projects.csv"), projects)

    qa = {
        "原始记录数": len(rows),
        "成功定位项目数": len(projects),
        "点位总数": len(points),
        "县级精度点位": sum(1 for p in points if p["定位精度"] == "县级"),
        "市级精度点位": sum(1 for p in points if p["定位精度"] == "市级"),
        "省级精度点位": sum(1 for p in points if p["定位精度"] == "省级"),
        "落在省界外的点位": outside,
        "无省面命中点位": nohit,
        "胡焕庸线以东点位": sum(1 for p in points if p["胡焕庸线以东"] == "是"),
    }
    json.dump({"summary": qa, "issues": bad},
              open(os.path.join(DATA, "qa_report.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(json.dumps(qa, ensure_ascii=False, indent=1))
    print("issues", len(bad))
    for b in bad[:10]:
        print(b)


if __name__ == "__main__":
    main()
