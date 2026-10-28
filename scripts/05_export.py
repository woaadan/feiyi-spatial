# -*- coding: utf-8 -*-
"""导出网页数据与 Excel 数据表。"""
import json, os, csv, collections
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(ROOT, "data")


def load(p):
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def main():
    pts = load(os.path.join(DATA, "points.csv"))
    proj = load(os.path.join(DATA, "projects.csv"))
    ad = json.load(open(os.path.join(DATA, "app_data.json"), encoding="utf-8"))

    # 紧凑点位（给地图/表格）
    compact = [[p["seq"], p["名称"], p["类别"], p["批次"], p["数据集"], p["省"],
                p["市"], p["县区"], float(p["经度"]), float(p["纬度"]),
                p["定位精度"], p["链接"], p["地点类型"], p["地点片段"]]
               for p in pts]
    json.dump(compact, open(os.path.join(OUT, "map_points.json"), "w", encoding="utf-8"),
              ensure_ascii=False, separators=(",", ":"))

    # ---------- Excel ----------
    wb = openpyxl.Workbook()
    HDR_FILL = PatternFill("solid", fgColor="1F3B57")
    HDR_FONT = Font(color="FFFFFF", bold=True, size=10, name="微软雅黑")
    BODY_FONT = Font(size=9, name="微软雅黑")
    thin = Side(style="thin", color="D9D9D9")

    def sheet(title, headers, rows, widths=None, freeze="A2"):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HDR_FILL; cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
                cell.border = Border(bottom=thin)
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        return ws

    ws = wb.active; ws.title = "说明"
    lines = [
        ["国家级非物质文化遗产代表性项目 · 空间数据集"],
        [""],
        ["数据来源", "原始名录（4234 条，含五批国家级名录与 2018 年第一批国家传统工艺振兴目录）"],
        ["空间定位", "按「申报地区或单位」逐段解析至省/市/县，坐标取行政区划驻地（DataV 行政区划数据）"],
        ["定位精度", "县级 2687 点 / 市级 1264 点 / 省级 390 点（省级为全省性项目或省级机构，落于省会）"],
        ["质量校验", "项目序号末 6 位行政区划码交叉核对；全部 4341 个点位经省界面几何包含检验，0 个越界"],
        ["点位口径", "一条记录可对应多个申报地区，故点位数（4341）大于记录数（4234）"],
        ["统计口径", "除特别说明外，空间统计以「点位」为单位；市域统计中直辖市与省域单位按省级归并"],
        [""],
        ["工作表", "项目主表 / 点位明细 / 省级汇总 / 市级汇总 / 类别统计 / 区位商 / 批次重心 / 空间自相关 / 省×类别矩阵 / 地理编码复核"],
    ]
    for r in lines:
        ws.append(r)
    ws["A1"].font = Font(size=14, bold=True, name="微软雅黑", color="1F3B57")
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(size=10, bold=True, name="微软雅黑")
        row[1].font = Font(size=10, name="微软雅黑")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    sheet("项目主表",
          ["序号", "项目编号", "名称", "类别", "批次", "类型", "数据集", "申报地区",
           "主点位", "省", "市", "县区", "县区代码", "经度", "纬度", "最高精度", "点位数", "详情页"],
          [[p["序号"], p["项目编号"], p["名称"], p["类别"], p["批次"], p["类型"], p["数据集"],
            p["申报地区"], p["主点位_地点"], p["省"], p["市"], p["县区"], p["县区代码"],
            float(p["经度"]), float(p["纬度"]), p["最高精度"], p["点位数"], p["链接"]] for p in proj],
          [6, 10, 22, 14, 14, 10, 16, 34, 22, 12, 14, 14, 10, 10, 9, 9, 7, 42])

    sheet("点位明细",
          ["序号", "名称", "类别", "批次", "数据集", "申报地区", "地点片段", "省", "省代码",
           "市", "县区", "县区代码", "经度", "纬度", "定位精度", "地点类型", "胡焕庸线以东"],
          [[p["seq"], p["名称"], p["类别"], p["批次"], p["数据集"], p["申报地区"], p["地点片段"],
            p["省"], p["省代码"], p["市"], p["县区"], p["县区代码"], float(p["经度"]),
            float(p["纬度"]), p["定位精度"], p["地点类型"], p["胡焕庸线以东"]] for p in pts],
          [6, 22, 16, 14, 16, 34, 22, 12, 9, 14, 14, 10, 10, 9, 9, 14, 11])

    prov = ad["provinces"]
    sheet("省级汇总",
          ["省级", "点位数", "占比", "国土面积(万km²)", "密度(个/万km²)", "类别数",
           "类别多样性", "首位类别", "胡焕庸线以东"],
          [[p["name"], p["count"], round(p["share"], 4), round(p["area"] / 10000, 2),
            round(p["density"], 2), p["n_cat"], round(p["shannon"], 3), p["top_cat"], p["east_hu"]]
           for p in prov],
          [16, 9, 9, 14, 15, 9, 11, 18, 13])

    sheet("市级汇总",
          ["市域", "省级", "点位数", "类别数", "类别多样性", "首位类别", "经度", "纬度"],
          [[c["name"], c["prov"], c["count"], c["n_cat"], round(c["shannon"], 3),
            c["top_cat"], round(c["lng"], 4), round(c["lat"], 4)] for c in ad["cities"]],
          [22, 14, 9, 9, 11, 18, 10, 9])

    sheet("类别统计",
          ["类别", "点位数", "覆盖省级数", "集中度CR3", "首位省级", "首位占比", "多样性"],
          [[c["name"], c["count"], c["n_prov"], round(ad["cr3"].get(c["name"], 0), 4),
            c["top_prov"], round(c["top_share"], 4), round(c["entropy"], 3)]
           for c in ad["categories"]],
          [24, 9, 11, 11, 16, 10, 9])

    cats = ad["matrix"]["cats"]
    lq = ad["lq"]
    sheet("区位商",
          ["省级"] + cats,
          [[nme] + [round(lq[nme].get(c, 0), 2) for c in cats] for nme in ad["matrix"]["provinces"]],
          [14] + [11] * len(cats))

    sheet("省×类别矩阵",
          ["省级"] + cats,
          [[nme] + [ad["matrix"]["values"][i][j] for j in range(len(cats))]
           for i, nme in enumerate(ad["matrix"]["provinces"])],
          [14] + [10] * len(cats))

    bs = ad["batch_stats"]
    sheet("批次重心",
          ["批次", "点位数", "重心经度", "重心纬度", "标准差椭圆长轴(km)", "短轴(km)",
           "方位角(°)", "平均半径(km)", "胡焕庸线以东占比"],
          [[b, bs[b]["count"], round(bs[b]["lng"], 4), round(bs[b]["lat"], 4),
            round(bs[b]["major"], 1), round(bs[b]["minor"], 1), round(bs[b]["angle"], 1),
            round(bs[b]["radius_km"], 1), round(bs[b]["east_share"], 4)]
           for b in ad["batch_order"] if b in bs],
          [16, 9, 11, 11, 20, 10, 10, 12, 18])

    sheet("空间自相关",
          ["省级", "点位数", "密度(个/万km²)", "局部Moran's I", "p值", "聚类类型", "z得分", "邻接数"],
          [[l["name"], next((p["count"] for p in prov if p["name"] == l["name"]), 0),
            round(l["density"], 2), round(l["I"], 4), round(l["p"], 3), l["type"],
            round(l["z"], 3), ad["weights_degree"].get(l["name"], 0)] for l in ad["lisa"]],
          [16, 9, 15, 13, 8, 10, 9, 9])

    review_p = os.path.join(DATA, "geo", "review.csv")
    rv = []
    if os.path.exists(review_p):
        rv = load(review_p)
    sheet("地理编码复核",
          ["序号", "项目名称", "申报地区", "片段", "问题", "解析结果", "项目序号中的码"],
          [[r["序号"], r["项目名称"], r["申报地区"], r["片段"], r["问题"], r["解析码"], r["序号码"]]
           for r in rv],
          [6, 20, 34, 22, 24, 26, 18])

    xp = os.path.join(ROOT, "国家级非遗空间数据集.xlsx")
    wb.save(xp)
    print("saved", xp, os.path.getsize(xp) / 1024, "KB")
    print("map_points", os.path.getsize(os.path.join(OUT, "map_points.json")) / 1024, "KB")


if __name__ == "__main__":
    main()
